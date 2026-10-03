"""Atomic document ingestion and transparent hybrid retrieval for a local corpus."""
import csv
import hashlib
import io
import json
import re
import sqlite3
import threading
import uuid
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import numpy as np

from .config import DEFAULT_SETTINGS, EMBED_MODEL


def now():
    return datetime.now(timezone.utc).isoformat()


def extract_document(name: str, content: bytes) -> list[dict]:
    """Keep page/section/row locators. Never execute markup or follow links."""
    suffix = Path(name).suffix.lower()
    if len(content) > 15 * 1024 * 1024:
        raise ValueError('Maximum document size is 15 MB.')
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(content))
        if reader.is_encrypted:
            raise ValueError('Encrypted PDFs are not supported. Upload an unlocked copy.')
        if len(reader.pages) > 300:
            raise ValueError('Maximum PDF length is 300 pages.')
        pages = [{'text': p.extract_text() or '', 'locator': f'Page {i + 1}'} for i, p in enumerate(reader.pages)]
        if not any(len(p['text'].strip()) > 30 for p in pages):
            raise ValueError('No usable text found. Scanned PDFs need OCR before upload; OCR is not enabled.')
        return pages
    if suffix == '.docx':
        from docx import Document
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(x.file_size for x in archive.infolist()) > 40 * 1024 * 1024:
                raise ValueError('Expanded DOCX exceeds 40 MB.')
        doc = Document(io.BytesIO(content))
        sections, text, heading = [], [], 'Document'
        for p in doc.paragraphs:
            if p.style and p.style.name.startswith('Heading'):
                if text:
                    sections.append({'text': '\n'.join(text), 'locator': heading})
                heading, text = p.text, [p.text]
            elif p.text.strip():
                text.append(p.text)
        if text:
            sections.append({'text': '\n'.join(text), 'locator': heading})
        for i, table in enumerate(doc.tables):
            rows = [' | '.join(c.text for c in r.cells) for r in table.rows]
            sections.append({'text': '\n'.join(rows), 'locator': f'Table {i + 1}'})
        return sections
    text = content.decode('utf-8-sig', errors='strict')
    if suffix in {'.html', '.htm'}:
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(text, 'html.parser')
        for tag in soup(['script', 'style', 'iframe', 'nav']):
            tag.decompose()
        return [{'text': soup.get_text('\n', strip=True), 'locator': 'HTML text'}]
    if suffix == '.csv':
        reader = csv.DictReader(io.StringIO(text))
        return [{'text': '; '.join(f'{k}: {v}' for k, v in row.items()), 'locator': f'Row {i + 2}'} for i, row in enumerate(reader)]
    if suffix not in {'.md', '.txt'}:
        raise ValueError('Supported formats: PDF with text, DOCX, TXT, Markdown, HTML, CSV.')
    sections, body, heading = [], [], 'Document'
    for line in text.splitlines():
        if re.match(r'^#{1,6}\s+', line):
            if body:
                sections.append({'text': '\n'.join(body), 'locator': heading})
            heading = re.sub(r'^#+\s*', '', line)
            body = [heading]
        else:
            body.append(line)
    if body:
        sections.append({'text': '\n'.join(body), 'locator': heading})
    return sections


def chunk_sections(sections, size=220, overlap=35):
    chunks = []
    for section in sections:
        words = section['text'].split()
        if not words:
            continue
        for start in range(0, len(words), size - overlap):
            body = ' '.join(words[start:start + size])
            chunks.append({'text': body, 'locator': section['locator']})
            if start + size >= len(words):
                break
    if not chunks:
        raise ValueError('This document contains no extractable text.')
    if len(chunks) > 1200:
        raise ValueError('Document exceeds 1,200 chunks. Split it into smaller files.')
    return chunks


STOP_WORDS = set('a an the is are am was were do does did what when where how why who which can could would should i you me we it they this that these those of for to on in with and or my your please tell about have has be at from much'.split())


def terms(text):
    return [t for t in re.findall(r'[\w]+', text.lower()) if len(t) > 1 and t not in STOP_WORDS][:32]


class KnowledgeStore:
    def __init__(self, directory: Path):
        directory.mkdir(parents=True, exist_ok=True)
        self.directory = directory
        self.lock = threading.RLock()
        self.db = sqlite3.connect(directory / 'studio.sqlite3', check_same_thread=False)
        self.db.row_factory = sqlite3.Row
        self.db.executescript('''
            PRAGMA journal_mode=WAL;
            PRAGMA foreign_keys=ON;
            CREATE TABLE IF NOT EXISTS documents(id TEXT PRIMARY KEY, name TEXT, sha TEXT UNIQUE,
                created TEXT, chunks INTEGER, chars INTEGER, embedding_model TEXT, demo INTEGER DEFAULT 0, raw BLOB);
            CREATE TABLE IF NOT EXISTS chunks(id TEXT PRIMARY KEY, document_id TEXT REFERENCES documents(id) ON DELETE CASCADE,
                text TEXT, locator TEXT, embedding BLOB);
            CREATE VIRTUAL TABLE IF NOT EXISTS chunks_fts USING fts5(id UNINDEXED, text, tokenize='porter unicode61');
            CREATE TABLE IF NOT EXISTS settings(key TEXT PRIMARY KEY, value TEXT);
            CREATE TABLE IF NOT EXISTS sessions(id TEXT PRIMARY KEY, created TEXT, ended TEXT, extraction TEXT);
            CREATE TABLE IF NOT EXISTS turns(id TEXT PRIMARY KEY, session_id TEXT REFERENCES sessions(id) ON DELETE CASCADE,
                created TEXT, user_text TEXT, answer TEXT, sources TEXT, metrics TEXT, status TEXT);
            CREATE TABLE IF NOT EXISTS rag_queries(id TEXT PRIMARY KEY, created TEXT, source TEXT,
                session_id TEXT, turn_id TEXT, query TEXT, elapsed_ms INTEGER, result_count INTEGER,
                results TEXT, error TEXT);
            CREATE INDEX IF NOT EXISTS rag_queries_created ON rag_queries(created DESC);
        ''')
        session_columns = {row['name'] for row in self.db.execute('PRAGMA table_info(sessions)')}
        if 'memory' not in session_columns:
            self.db.execute('ALTER TABLE sessions ADD COLUMN memory TEXT')
        self.db.commit()

    def settings(self):
        with self.lock:
            settings = dict(DEFAULT_SETTINGS)
            for row in self.db.execute('SELECT * FROM settings'):
                settings[row['key']] = json.loads(row['value'])
            return settings

    def save_settings(self, settings):
        with self.lock, self.db:
            for key, value in settings.items():
                self.db.execute('INSERT OR REPLACE INTO settings VALUES (?,?)', (key, json.dumps(value)))

    def documents(self):
        with self.lock:
            return [dict(x) for x in self.db.execute('SELECT id,name,sha,created,chunks,chars,embedding_model,demo FROM documents ORDER BY created DESC')]

    def find_sha(self, sha):
        return next((x for x in self.documents() if x['sha'] == sha), None)

    def insert_document(self, name, raw, chunks, vectors, demo=False):
        if len(chunks) != len(vectors):
            raise ValueError('Embedding count does not match chunk count.')
        sha = hashlib.sha256(raw).hexdigest()
        doc_id = uuid.uuid4().hex
        with self.lock, self.db:
            duplicate = self.db.execute('SELECT id FROM documents WHERE sha=?', (sha,)).fetchone()
            if duplicate:
                return duplicate['id']
            self.db.execute('INSERT INTO documents VALUES (?,?,?,?,?,?,?,?,?)',
                (doc_id, Path(name).name[:200], sha, now(), len(chunks), sum(len(c['text']) for c in chunks), EMBED_MODEL, int(demo), raw))
            for chunk, vector in zip(chunks, vectors):
                chunk_id = uuid.uuid4().hex[:12]
                embedding = np.asarray(vector, dtype=np.float32)
                if not np.isfinite(embedding).all() or not np.linalg.norm(embedding):
                    raise ValueError('Invalid embedding vector.')
                embedding /= np.linalg.norm(embedding)
                self.db.execute('INSERT INTO chunks VALUES (?,?,?,?,?)', (chunk_id, doc_id, chunk['text'], chunk['locator'], embedding.tobytes()))
                self.db.execute('INSERT INTO chunks_fts VALUES (?,?)', (chunk_id, chunk['text']))
        return doc_id

    def delete_document(self, doc_id):
        with self.lock, self.db:
            exists = self.db.execute('SELECT 1 FROM documents WHERE id=?', (doc_id,)).fetchone()
            if not exists:
                return False
            self.db.execute('DELETE FROM chunks_fts WHERE id IN (SELECT id FROM chunks WHERE document_id=?)', (doc_id,))
            result = self.db.execute('DELETE FROM documents WHERE id=?', (doc_id,))
            return result.rowcount > 0

    def document_detail(self, doc_id):
        with self.lock:
            doc = self.db.execute('SELECT id,name,created,chunks,chars,demo,embedding_model FROM documents WHERE id=?', (doc_id,)).fetchone()
            if not doc:
                return None
            return {**dict(doc), 'sections': [dict(x) for x in self.db.execute('SELECT id,text,locator FROM chunks WHERE document_id=?', (doc_id,))]}

    def search(self, query, vector, limit=4, threshold=0.45):
        """Rank lexical BM25 + cosine using RRF; expose scores, never call them probabilities."""
        with self.lock:
            rows = [dict(r) for r in self.db.execute('''SELECT c.id,c.text,c.locator,c.embedding,
                d.name,d.id as document_id,d.demo FROM chunks c JOIN documents d ON d.id=c.document_id
                WHERE d.embedding_model=?''', (EMBED_MODEL,))]
            tokens = terms(query)
            fts = ' OR '.join('"' + t.replace('"', '') + '"' for t in tokens)
            lexical = self.db.execute('SELECT id,bm25(chunks_fts) as score FROM chunks_fts WHERE chunks_fts MATCH ? ORDER BY score LIMIT 24', (fts,)).fetchall() if fts else []
        if not rows:
            return []
        q = np.asarray(vector, dtype=np.float32)
        q /= max(float(np.linalg.norm(q)), 1e-8)
        lexical_rank = {r['id']: rank for rank, r in enumerate(lexical)}
        for row in rows:
            v = np.frombuffer(row.pop('embedding'), dtype=np.float32)
            if v.shape != q.shape:
                raise ValueError('Embedding dimensions changed. Re-ingest documents with the current model.')
            row['similarity'] = float(np.dot(v, q))
            row['keyword_match'] = row['id'] in lexical_rank
        rows.sort(key=lambda r: r['similarity'], reverse=True)
        for rank, row in enumerate(rows):
            row['score'] = 1 / (60 + rank + 1)
            if row['id'] in lexical_rank:
                row['score'] += 1 / (60 + lexical_rank[row['id']] + 1)
        rows = [r for r in rows if r['similarity'] >= threshold]
        rows.sort(key=lambda r: r['score'], reverse=True)
        return rows[:limit]

    def create_session(self):
        sid = uuid.uuid4().hex
        with self.lock, self.db:
            self.db.execute(
                'INSERT INTO sessions(id,created,ended,extraction,memory) VALUES (?,?,NULL,NULL,NULL)',
                (sid, now()),
            )
        return sid

    def end_session(self, sid):
        with self.lock, self.db:
            self.db.execute('UPDATE sessions SET ended=? WHERE id=?', (now(), sid))

    def add_turn(self, sid, tid, user, answer, sources, metrics, status='complete'):
        with self.lock, self.db:
            self.db.execute('INSERT OR REPLACE INTO turns VALUES (?,?,?,?,?,?,?,?)',
                (tid, sid, now(), user, answer, json.dumps(sources), json.dumps(metrics), status))

    def add_rag_query(self, query, source, elapsed_ms, results=None, error=None,
                      session_id=None, turn_id=None):
        results = results or []
        evidence = [{key: item.get(key) for key in
                     ('id', 'document_id', 'name', 'locator', 'similarity', 'score')}
                    for item in results]
        with self.lock, self.db:
            self.db.execute('INSERT INTO rag_queries VALUES (?,?,?,?,?,?,?,?,?,?)', (
                uuid.uuid4().hex, now(), source, session_id, turn_id, query,
                int(elapsed_ms), len(results), json.dumps(evidence), error,
            ))

    def rag_queries(self, limit=200):
        limit = max(1, min(int(limit), 1000))
        with self.lock:
            rows = [dict(row) for row in self.db.execute(
                'SELECT * FROM rag_queries ORDER BY created DESC LIMIT ?', (limit,))]
        for row in rows:
            row['results'] = json.loads(row['results'])
        return rows

    def sessions(self):
        with self.lock:
            return [dict(r) for r in self.db.execute('''SELECT s.*, count(t.id) as turns,
                substr(min(t.user_text),1,80) as preview FROM sessions s LEFT JOIN turns t ON s.id=t.session_id
                GROUP BY s.id ORDER BY s.created DESC LIMIT 100''')]

    def session(self, sid):
        with self.lock:
            row = self.db.execute('SELECT * FROM sessions WHERE id=?', (sid,)).fetchone()
            if not row:
                return None
            result = dict(row)
            result['turns'] = [dict(t) for t in self.db.execute('SELECT * FROM turns WHERE session_id=? ORDER BY created', (sid,))]
        for t in result['turns']:
            t['sources'] = json.loads(t['sources'])
            t['metrics'] = json.loads(t['metrics'])
        result['extraction'] = json.loads(result['extraction']) if result['extraction'] else None
        result['memory'] = json.loads(result['memory']) if result.get('memory') else None
        return result

    def save_extraction(self, sid, extraction):
        with self.lock, self.db:
            self.db.execute('UPDATE sessions SET extraction=? WHERE id=?', (json.dumps(extraction), sid))

    def save_memory(self, sid, memory):
        with self.lock, self.db:
            self.db.execute('UPDATE sessions SET memory=? WHERE id=?', (json.dumps(memory), sid))
